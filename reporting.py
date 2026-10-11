"""Database-backed circulation analytics and portable report exports."""
import csv
import io
import math
from datetime import datetime, date, time, timedelta, timezone
from pathlib import Path
from xml.sax.saxutils import escape

from sqlalchemy import or_, func
from sqlalchemy.orm import joinedload
from models import Book, BorrowRecord, BorrowStatus, Fine, User, UserRole

CAMPUS = "EVSU Ormoc City Campus"
LOCAL_TZ = timezone(timedelta(hours=8))
ROOT = Path(__file__).resolve().parent


def local_day(value):
    return value.replace(tzinfo=timezone.utc).astimezone(LOCAL_TZ).date()


def dashboard_stats(db):
    books = db.query(Book).filter(Book.is_archived.is_(False))
    return {
        "total_books": books.with_entities(func.coalesce(func.sum(Book.total_copies), 0)).scalar(),
        "catalog_titles": books.count(),
        "active_borrows": db.query(BorrowRecord).filter(BorrowRecord.status.in_(
            [BorrowStatus.CHECKED_OUT, BorrowStatus.OVERDUE])).count(),
        "pending_requests": db.query(BorrowRecord).filter(BorrowRecord.status == BorrowStatus.PENDING).count(),
        "registered_patrons": db.query(User).filter(User.role.in_([UserRole.STUDENT, UserRole.FACULTY])).count(),
        "uncollected_fines": float(db.query(func.coalesce(func.sum(Fine.amount), 0)).filter(Fine.is_paid.is_(False)).scalar()),
    }


def circulation_report(db, start: date, end: date, department=None, course=None, role="student", sort="date", direction="asc"):
    start_utc = datetime.combine(start, time.min, LOCAL_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    end_utc = datetime.combine(end + timedelta(days=1), time.min, LOCAL_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    query = db.query(BorrowRecord).join(User).options(joinedload(BorrowRecord.user), joinedload(BorrowRecord.book)).filter(or_(
        (BorrowRecord.issue_date >= start_utc) & (BorrowRecord.issue_date < end_utc),
        (BorrowRecord.return_date >= start_utc) & (BorrowRecord.return_date < end_utc),
    ))
    if role != "all":
        query = query.filter(User.role == UserRole(role))
    if department:
        query = query.filter(User.department == department)
    if course:
        query = query.filter(User.course == course)
    daily = {str(start + timedelta(days=n)): {"date": str(start + timedelta(days=n)), "borrowed_items": 0,
        "returned_items": 0, "borrowers": set(), "returners": set()} for n in range((end - start).days + 1)}
    rows, borrowers, returners = [], set(), set()
    for record in query.all():
        for action, stamp in [("Borrowed", record.issue_date), ("Returned", record.return_date)]:
            if action == "Borrowed" and record.status not in (BorrowStatus.CHECKED_OUT, BorrowStatus.OVERDUE, BorrowStatus.RETURNED):
                continue
            if not stamp or not start_utc <= stamp < end_utc:
                continue
            day = str(local_day(stamp))
            count_key, people_key = ("borrowed_items", "borrowers") if action == "Borrowed" else ("returned_items", "returners")
            daily[day][count_key] += 1
            daily[day][people_key].add(record.user_id)
            (borrowers if action == "Borrowed" else returners).add(record.user_id)
            rows.append({"date": day, "time": stamp.replace(tzinfo=timezone.utc).astimezone(LOCAL_TZ).strftime("%H:%M"),
                "name": record.user.full_name, "institutional_id": record.user.username,
                "department": record.user.department or "Not recorded", "course": record.user.course or "Not recorded",
                "book": record.book.title if record.book else "Removed book", "action": action, "reference": record.req_id})
    rows.sort(key=lambda r: (r[sort].casefold(), r["date"], r["time"], r["institutional_id"], r["reference"]), reverse=direction == "desc")
    series = [{**item, "borrowers": len(item["borrowers"]), "returners": len(item["returners"])} for item in daily.values()]
    return {"campus": CAMPUS, "start": str(start), "end": str(end), "timezone": "Asia/Manila (UTC+08:00)",
        "filters": {"department": department or "All", "course": course or "All", "role": role, "sort": sort, "direction": direction},
        "generated_at": datetime.now(LOCAL_TZ).isoformat(),
        "summary": {"borrowed_items": sum(d["borrowed_items"] for d in series), "returned_items": sum(d["returned_items"] for d in series),
                    "unique_borrowers": len(borrowers), "unique_returners": len(returners)},
        "daily": series, "rows": rows}


def font(size, bold=False):
    from PIL import ImageFont
    candidates = [ROOT / "static" / ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"),
        Path("C:/Windows/Fonts") / ("arialbd.ttf" if bold else "arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu") / ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf")]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default(size=size)


def chart_png(report):
    from PIL import Image, ImageDraw
    image = Image.new("RGB", (1400, 520), "white")
    draw = ImageDraw.Draw(image)
    draw.text((65, 20), "Daily people borrowing and returning", fill="#252129", font=font(28, True))
    series = report["daily"]
    high = max(1, max((max(d["borrowers"], d["returners"]) for d in series), default=0))
    step = max(1, math.ceil(high / 4))
    high = step * 4
    left, top, right, bottom = 80, 110, 1320, 410
    for n in range(5):
        y = bottom - n * (bottom-top) / 4
        draw.line((left, y, right, y), fill="#e6e3e3", width=2)
        draw.text((25, y-12), str(n*step), fill="#666666", font=font(20))
    for key, color, title, lx in [("borrowers", "#730d24", "Borrowers", 80), ("returners", "#237966", "Returners", 290)]:
        draw.line((lx, 78, lx+38, 78), fill=color, width=5)
        draw.text((lx+50, 62), title, fill=color, font=font(21))
        points = [(left + i * (right-left)/max(1, len(series)-1), bottom - d[key] / high * (bottom-top)) for i, d in enumerate(series)]
        if len(points)>1:
            draw.line(points, fill=color, width=4)
        for x,y in points:
            draw.ellipse((x-4,y-4,x+4,y+4), fill=color)
    for i in sorted(set([0, len(series)-1] + list(range(0, len(series), max(1, len(series)//5))))):
        x = left + i * (right-left)/max(1, len(series)-1)
        draw.text((x, 430), series[i]["date"][5:], fill="#666666", font=font(18), anchor="mt")
    draw.text((80, 478), "Unique people per day; a person borrowing multiple items counts once. Dates: Philippine time.", fill="#666666", font=font(18))
    data=io.BytesIO(); image.save(data, format="PNG"); return data.getvalue()


COLUMNS = [("date", "Date"), ("time", "Time"), ("name", "Name"), ("institutional_id", "ID"), ("department", "Department"),
           ("course", "Course"), ("book", "Book"), ("action", "Activity"), ("reference", "Reference")]


def export_report(report, format):
    if format == "csv":
        output=io.StringIO(newline=""); writer=csv.writer(output)
        def safe(value):
            value = str(value)
            return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value
        writer.writerow(["Library circulation report", CAMPUS])
        writer.writerow(["From", report["start"], "To", report["end"], "Timezone", report["timezone"]])
        writer.writerow(["Audience", report["filters"]["role"], "Department", safe(report["filters"]["department"]), "Course", safe(report["filters"]["course"])])
        writer.writerow([])
        writer.writerow(["Period totals", "Count"])
        for key, value in report["summary"].items():
            writer.writerow([key.replace("_", " ").title(), value])
        writer.writerow([])
        writer.writerow(["Daily totals (unique people per day)"])
        writer.writerow(["Date", "Borrowers", "Returners", "Items borrowed", "Items returned"])
        for day in report["daily"]:
            writer.writerow([day[k] for k in ("date", "borrowers", "returners", "borrowed_items", "returned_items")])
        writer.writerow([])
        writer.writerow(["Activity details"])
        writer.writerow([label for _,label in COLUMNS])
        for row in report["rows"]:
            # Prevent values supplied by patrons from becoming spreadsheet formulas.
            writer.writerow([safe(row[k]) for k,_ in COLUMNS])
        return output.getvalue().encode("utf-8-sig")
    filters=report["filters"]
    title="Library circulation report"
    subtitle=f"{CAMPUS} | {report['start']} to {report['end']} | Philippine time"
    scope=f"Audience: {filters['role']} | Department: {filters['department']} | Course: {filters['course']}"
    summary=" | ".join(f"{key.replace('_', ' ').title()}: {value}" for key,value in report["summary"].items())
    chart=chart_png(report)
    output=io.BytesIO()
    if format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, LongTable, TableStyle
        styles=getSampleStyleSheet(); small=ParagraphStyle("cell", fontSize=7, leading=10)
        content=[Paragraph(title, styles["Title"]), Paragraph(escape(subtitle), styles["Normal"]),
                 Paragraph(escape(scope), styles["Normal"]), Spacer(1,12), Paragraph(escape(summary), styles["Normal"]),
                 Image(io.BytesIO(chart), width=750, height=279),
                 Paragraph("Activity details (actual issue and return dates)", styles["Heading2"])]
        rows=[[Paragraph(escape(label),small) for _,label in COLUMNS]]
        rows += [[Paragraph(escape(str(row[k])),small) for k,_ in COLUMNS] for row in report["rows"]]
        if report["rows"]:
            table=LongTable(rows, colWidths=[55,32,95,68,85,100,130,55,70], repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#f3e7e9")),("VALIGN",(0,0),(-1,-1),"TOP"),
                ("GRID",(0,0),(-1,-1),.3,colors.HexColor("#dddddd")),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
            content.append(table)
        else: content.append(Paragraph("No recorded circulation activity matches these filters.",styles["Normal"]))
        def footer(canvas, doc):
            canvas.setFont("Helvetica",8); canvas.drawString(32,18,f"{CAMPUS} | Generated {report['generated_at'][:19]} | Page {doc.page}")
        SimpleDocTemplate(output,pagesize=landscape(A4),rightMargin=26,leftMargin=26,topMargin=26,bottomMargin=32).build(content,onFirstPage=footer,onLaterPages=footer)
    else:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.section import WD_ORIENT
        doc=Document(); section=doc.sections[0]; section.orientation=WD_ORIENT.LANDSCAPE
        section.page_width=Inches(11.7); section.page_height=Inches(8.3)
        section.left_margin=section.right_margin=Inches(.45)
        doc.add_heading(title,0); doc.add_paragraph(subtitle); doc.add_paragraph(scope); doc.add_paragraph(summary)
        doc.add_picture(io.BytesIO(chart),width=Inches(10.5)); doc.add_heading("Activity details",1)
        table=doc.add_table(rows=1,cols=len(COLUMNS)); table.style="Light Shading Accent 1"
        from docx.oxml import OxmlElement
        repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for cell,(_,label) in zip(table.rows[0].cells,COLUMNS): cell.text=label
        for row in report["rows"]:
            for cell,(key,_) in zip(table.add_row().cells,COLUMNS): cell.text=str(row[key])
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs: run.font.size=Pt(8)
        if not report["rows"]: doc.add_paragraph("No recorded circulation activity matches these filters.")
        section.footer.paragraphs[0].text=f"{CAMPUS} | Generated {report['generated_at'][:19]}"
        doc.save(output)
    return output.getvalue()


def library_id_png(user):
    from PIL import Image, ImageDraw
    import qrcode
    image=Image.new("RGB",(900,1400),"#ffffff"); draw=ImageDraw.Draw(image)
    wine, gold, ink="#730d24", "#d8b366", "#252129"
    draw.polygon([(0,0),(335,0),(0,335)],fill=wine)
    draw.polygon([(900,0),(665,0),(900,235)],fill=wine)
    draw.line([(0,350),(350,0)], fill=gold,width=14)
    draw.line([(900,250),(650,0)],fill=gold,width=14)
    logo_path=ROOT/"mobile-native"/"assets"/"evsu-smartlib-logo.jpg"
    if logo_path.exists():
        logo=Image.open(logo_path).convert("RGB").resize((150,150)); image.paste(logo,(375,48))
    def centered(text,y,size=30,bold=False,color=ink,max_width=780):
        words=str(text or "Not recorded").split(); lines=[]; line=""
        for word in words:
            test=(line+" "+word).strip()
            if draw.textlength(test,font=font(size,bold))>max_width and line: lines.append(line); line=word
            else: line=test
        lines.append(line)
        for line in lines:
            while draw.textlength(line,font=font(size,bold))>max_width and size>12: size-=1
            draw.text((450,y),line,font=font(size,bold),fill=color,anchor="mt"); y+=size+10
        return y
    centered("EASTERN VISAYAS STATE UNIVERSITY",226,26,True)
    centered("ORMOC CITY CAMPUS",269,24,True,wine)
    centered("SMARTLIB  /  LIBRARY MEMBERSHIP",320,20,False)
    payload=f"EVSU-CARD-2026::{user.username}::{user.role.value}"
    qr=qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M,box_size=12,border=4)
    qr.add_data(payload); qr.make(fit=True)
    qr.box_size = max(1, 660 // (qr.modules_count + 8))
    qr_image=qr.make_image(fill_color="black",back_color="white").convert("RGB")
    # Integer module scaling preserves a four-module white quiet zone.
    side=qr_image.width; image.paste(qr_image,((900-side)//2,390))
    draw.rounded_rectangle((100,375,800,1070),radius=24,outline=gold,width=4)
    name_y=1110
    name_y=centered(user.full_name,name_y,32,True)
    name_y=centered(user.username,name_y+2,26,True,wine)
    centered(f"{user.role.value.title()} | {'Active member' if user.is_active else 'Inactive member'}",name_y+4,19)
    draw.rectangle((0,1250,900,1400),fill=wine)
    y=centered(user.department or "Ormoc City Campus",1270,22,True,"#ffffff")
    centered(user.course or "Library member",y+2,20,False,"#ffffff")
    output=io.BytesIO(); image.save(output,format="PNG"); return output.getvalue()
