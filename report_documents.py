"""Report documents on the retained EVSU letterhead, with growing table rows."""
from copy import deepcopy
import io
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape

from lxml import etree

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "templates" / "reports" / "template.docx"
LETTERHEAD = ROOT / "templates" / "reports" / "letterhead.jpeg"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def clean_text(value):
    # OOXML rejects C0 characters and lone surrogates; preserve paragraphs/tabs.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", str(value if value is not None else ""))


def _template_geometry(paper="long"):
    if paper not in {"long", "a4"}:
        raise ValueError("Paper must be a4 or long")
    with ZipFile(TEMPLATE) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    size = root.find(".//w:sectPr/w:pgSz", NS)
    margins = root.find(".//w:sectPr/w:pgMar", NS)
    # Long bond preserves native artwork; A4 scales the complete artwork to fit.
    width, height = (8.5 * 72, 13 * 72) if paper == "long" else tuple(
        int(size.get(f"{{{NS['w']}}}{key}")) / 20 for key in ("w", "h"))
    left, right = [int(margins.get(f"{{{NS['w']}}}{key}")) / 20 for key in ("left", "right")]
    return width, height, left, right


def _fit_header(payload, width, height):
    """Fit the source's oversized background without cropping or distortion."""
    root = etree.fromstring(payload)
    changed = False
    for anchor in root.findall(".//wp:anchor", NS):
        if anchor.get("behindDoc") != "1":
            continue
        extent = anchor.find("wp:extent", NS)
        if extent is None:
            continue
        cx, cy = int(extent.get("cx")), int(extent.get("cy"))
        scale = min(1.0, width * 12700 / cx, height * 12700 / cy)
        if scale == 1.0:
            continue
        fitted_x, fitted_y = round(cx * scale), round(cy * scale)
        extent.set("cx", str(fitted_x)); extent.set("cy", str(fitted_y))
        for child in anchor.findall(".//a:xfrm/a:ext", NS):
            child.set("cx", str(fitted_x)); child.set("cy", str(fitted_y))
        for axis, offset in (("H", round((width * 12700 - fitted_x) / 2)), ("V", 0)):
            position = anchor.find(f"wp:position{axis}", NS)
            if position is not None:
                position.set("relativeFrom", "page")
                for child in list(position): position.remove(child)
                etree.SubElement(position, f"{{{NS['wp']}}}posOffset").text = str(offset)
        changed = True
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True) if changed else payload


def _merge_parts(original, generated, overrides):
    """Keep opaque template parts intact instead of round-tripping them."""
    with ZipFile(original) as source, ZipFile(io.BytesIO(generated)) as built:
        parts = {name: source.read(name) for name in source.namelist()}
        parts["word/document.xml"] = built.read("word/document.xml")
        for name, identity in (("word/_rels/document.xml.rels", "Id"), ("[Content_Types].xml", None)):
            source_root, new_root = etree.fromstring(parts[name]), etree.fromstring(built.read(name))
            def key(element):
                return element.get(identity) if identity else (element.tag, element.get("PartName") or element.get("Extension"))
            known = {key(element) for element in source_root}
            added = False
            for element in new_root:
                if key(element) not in known:
                    source_root.append(deepcopy(element)); added = True
            if added:
                parts[name] = etree.tostring(source_root, xml_declaration=True, encoding="UTF-8", standalone=True)
        for name in built.namelist():
            if name not in parts: parts[name] = built.read(name)
        parts.update(overrides)
        out = io.BytesIO()
        with ZipFile(out, "w", ZIP_DEFLATED) as target:
            for name, data in parts.items(): target.writestr(name, data)
    return out.getvalue()


def _docx(title, metadata, sections, chart, paper):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(TEMPLATE)
    page_width, page_height, left, right = _template_geometry(paper)
    available = page_width - left - right
    # Source margins would collide with its graphic. Reserve the full artwork.
    for section in doc.sections:
        section.page_width = Pt(page_width); section.page_height = Pt(page_height)
        section.top_margin = Pt(110); section.bottom_margin = Pt(78)
        section.different_first_page_header_footer = False
    doc.settings.odd_and_even_pages_header_footer = False
    first_slot = doc.paragraphs[0] if doc.paragraphs else None
    for element in list(doc.element.body):
        if element.tag != qn("w:sectPr") and (first_slot is None or element is not first_slot._p):
            doc.element.body.remove(element)
    title_slot = [first_slot] if first_slot is not None else []

    def paragraph(text, size=11, bold=False, keep=False):
        p = title_slot.pop() if title_slot else doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = keep
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(clean_text(text)); run.font.name = "Times New Roman"
        run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = RGBColor(0, 0, 0)
        return p
    paragraph(title, 16, True, True)
    for text in metadata: paragraph(text, 10)
    if chart:
        doc.add_picture(io.BytesIO(chart), width=Pt(available))
    for section in sections:
        paragraph(section["title"], 12, True, True)
        columns, rows = section["columns"], section["rows"]
        if not rows:
            paragraph(section.get("empty", "No recorded activity matches these filters."), 10)
            continue
        table = doc.add_table(rows=1, cols=len(columns))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
        weights = section.get("widths") or [1] * len(columns)
        widths = [available * weight / sum(weights) for weight in weights]
        for column, width in zip(table.columns, widths): column.width = Pt(width)
        properties = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}"); border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4"); border.set(qn("w:color"), "D9D9D9"); borders.append(border)
        properties.append(borders)
        for values in rows: table.add_row()
        for index, (row, values) in enumerate(zip(table.rows, [columns] + rows)):
            # No exact/minimum row height or no-wrap: Word expands and splits.
            row.height = None
            if index == 0:
                repeat = OxmlElement("w:tblHeader"); row._tr.get_or_add_trPr().append(repeat)
            for cell, value, width in zip(row.cells, values, widths):
                cell.width = Pt(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = clean_text(value)
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    edge_element = OxmlElement(f"w:{edge}"); edge_element.set(qn("w:w"), "70")
                    edge_element.set(qn("w:type"), "dxa"); margins.append(edge_element)
                cell._tc.get_or_add_tcPr().append(margins)
                if index == 0:
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "EEEEEE"); cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.keep_with_next = index == 0
                    for run in p.runs:
                        run.font.name = "Times New Roman"; run.font.size = Pt(9)
                        run.bold = index == 0; run.font.color.rgb = RGBColor(0, 0, 0)
        paragraph("", 3)
    generated = io.BytesIO(); doc.save(generated)
    with ZipFile(TEMPLATE) as source:
        overrides = {name: _fit_header(source.read(name), page_width, page_height)
                     for name in source.namelist() if re.fullmatch(r"word/header\d+\.xml", name)}
    return _merge_parts(TEMPLATE, generated.getvalue(), overrides)


def _pdf(title, metadata, sections, chart, paper):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Image, CondPageBreak
    from PIL import Image as PILImage
    font, bold = "Times-Roman", "Times-Bold"
    fonts = Path("C:/Windows/Fonts")
    if (fonts / "times.ttf").exists() and (fonts / "timesbd.ttf").exists():
        if "ReportTimes" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("ReportTimes", str(fonts / "times.ttf")))
            pdfmetrics.registerFont(TTFont("ReportTimesBold", str(fonts / "timesbd.ttf")))
        font, bold = "ReportTimes", "ReportTimesBold"
    width, height, left, right = _template_geometry(paper); available = width - left - right
    normal = ParagraphStyle("report-body", fontName=font, fontSize=10, leading=13, spaceAfter=6)
    heading = ParagraphStyle("report-heading", parent=normal, fontName=bold, fontSize=12, leading=15, spaceBefore=9, spaceAfter=6)
    title_style = ParagraphStyle("report-title", parent=heading, fontSize=16, leading=19, spaceBefore=0)
    cell_style = ParagraphStyle("report-cell", parent=normal, fontSize=9, leading=11, spaceAfter=0, splitLongWords=True)
    header_style = ParagraphStyle("report-column", parent=cell_style, fontName=bold)

    class ReportTable(LongTable):
        def _splitRows(self, availHeight, doInRowSplit=0):
            # ReportLab retries the continuation in the current page's leftover
            # space. Splitting an ordinary row there repeats the header at the
            # bottom and separates one record across pages. Only a record too
            # tall for a fresh frame needs the within-row fallback.
            if doInRowSplit:
                repeated = self.repeatRows if isinstance(self.repeatRows, int) else max(self.repeatRows, default=-1) + 1
                frame = getattr(self, "_frame", None)
                if frame is not None and repeated < len(self._rowHeights):
                    first_record_height = sum(self._rowHeights[:repeated + 1])
                    if first_record_height <= frame._aH:
                        return []
            return super()._splitRows(availHeight, doInRowSplit=doInRowSplit)

    def para(value, style=normal): return Paragraph(escape(clean_text(value)).replace("\n", "<br/>"), style)
    content = [para(title, title_style)] + [para(line) for line in metadata]
    if chart:
        with PILImage.open(io.BytesIO(chart)) as picture: ratio = picture.height / picture.width
        content.extend([Image(io.BytesIO(chart), width=available, height=available * ratio), Spacer(1, 8)])
    for section in sections:
        content.append(CondPageBreak(96))
        content.append(para(section["title"], heading))
        if not section["rows"]:
            content.append(para(section.get("empty", "No recorded activity matches these filters.")))
            continue
        weights = section.get("widths") or [1] * len(section["columns"])
        rows = [[para(value, header_style) for value in section["columns"]]]
        rows += [[para(value, cell_style) for value in row] for row in section["rows"]]
        table = ReportTable(rows, colWidths=[available * weight / sum(weights) for weight in weights], repeatRows=1,
                          splitByRow=1, splitInRow=1, hAlign="LEFT")
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#d9d9d9")),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        content.extend([table, Spacer(1, 8)])
    def page(canvas, document):
        canvas.saveState()
        # Original drawing extent in the supplied DOCX, anchored at page top left.
        artwork_width, artwork_height = 7756525 / 12700, 11863070 / 12700
        scale = min(1.0, width / artwork_width, height / artwork_height)
        image_width, image_height = artwork_width * scale, artwork_height * scale
        x = (width - image_width) / 2 if paper == "a4" else 0
        canvas.drawImage(ImageReader(str(LETTERHEAD)), x, height - image_height,
                         image_width, image_height)
        canvas.restoreState()
    output = io.BytesIO()
    SimpleDocTemplate(output, pagesize=(width, height), leftMargin=left, rightMargin=right,
                      topMargin=110, bottomMargin=78, title=clean_text(title), author="EVSU Ormoc City Campus").build(
                          content, onFirstPage=page, onLaterPages=page)
    return output.getvalue()


def export_document(title, metadata, sections, format, chart=None, paper="long"):
    if format == "docx": return _docx(title, metadata, sections, chart, paper)
    if format == "pdf": return _pdf(title, metadata, sections, chart, paper)
    raise ValueError("Document format must be pdf or docx")
